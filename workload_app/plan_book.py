"""Projects and submissions, out to Excel and back in (Ahmed, 2026-10-12).

A manager with a lot to change at once -- every submission's planned day
after a client meeting, a dozen budgets -- does it faster in Excel.  So the
plan goes out as a workbook with two sheets:

* **Projects** -- job number, name, status, budget, start, end, notes;
* **Submissions** -- each deliverable's newest sending: revision, what it is
  sent for, the day it is planned to go, the day it went, the day it came
  back with its code and reason, the day its work starts and who prepares it.

Brought back through Bring in data, it is read, compared with what the app
holds now, and every difference is listed for the manager to see before
anything is saved.  Only what changed is written.  Nothing is ever deleted
from Excel: a row taken out of the sheet is simply left as it is, and a new
job number with a name and a budget is a new project.  A submission is found
by its **Ref**, the grey first column, so its name can be edited elsewhere
without losing track of it.
"""

from __future__ import annotations

import datetime as _dt
import io
from typing import Any, Dict, List, Optional, Sequence, Tuple

from . import revisions
from .model import as_date, as_number, as_text

TITLE = "Selecao+ projects and submissions"

PROJECT_COLUMNS = [
    ("number", "Job no."), ("name", "Project name"), ("status", "Status"),
    ("budget_mm", "Budget (MM)"), ("start", "Start"), ("end", "End"),
    ("notes", "Notes"),
]
SUBMISSION_COLUMNS = [
    ("ref", "Ref"), ("job", "Job no."), ("deliverable", "Deliverable"),
    ("rev", "Revision"), ("purpose", "Sent for"), ("planned", "Planned to go"),
    ("submitted", "Submitted"), ("returned", "Came back"),
    ("code", "Client code (A/B/C)"), ("reason", "Reason or comments"),
    ("starts", "Work starts"), ("people", "Who prepares it"),
]
#: Columns read back; the rest are there to read.
PROJECT_FIELDS = ("name", "status", "budget_mm", "start", "end", "notes")
ISSUE_FIELDS = ("rev", "purpose", "planned", "submitted", "returned", "code", "reason")
DATE_FIELDS = {"start", "end", "planned", "submitted", "returned", "starts"}
LABELS = {**dict(PROJECT_COLUMNS), **dict(SUBMISSION_COLUMNS),
          "name": "Project name", "status": "Status"}


class PlanBookError(ValueError):
    """The workbook is ours but cannot be read."""


# --------------------------------------------------------------------------
# out
# --------------------------------------------------------------------------

def write(*, projects: Sequence[Dict[str, Any]], items: Sequence[Dict[str, Any]],
          starts: Dict[int, Optional[str]], statuses: Sequence[str],
          unit_name: str = "", now: Optional[_dt.datetime] = None) -> bytes:
    """The workbook, as ``.xlsx`` bytes."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    now = now or _dt.datetime.now()
    book = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F2A44")
    grey = PatternFill("solid", fgColor="E6E8EC")
    grey_font = Font(color="5B6472")

    about = book.active
    about.title = "Read me"
    lines = [
        (TITLE, ""),
        ("Unit", unit_name),
        ("Written out", now.strftime("%d/%m/%Y %H:%M")),
        ("", ""),
        ("How to use it", "Change what you need on Projects and Submissions, save, "
                          "and bring the file back in on Bring in data."),
        ("", "You see every change before anything is saved, and only what changed is written."),
        ("", "Dates are day first: 25/12/2026."),
        ("", "Grey columns are there to read: Ref finds the submission, so leave it as it is."),
        ("", "A new row on Projects with a job number, a name and a budget becomes a new project."),
        ("", "Taking a row out does not remove anything from the app."),
        ("", "Who prepares it: names as on the team, separated by commas."),
        ("", "Client code: A accepted, B with comments, C revise and resubmit."),
    ]
    for line in lines:
        about.append(list(line))
    about["A1"].font = Font(bold=True, size=14)
    about.column_dimensions["A"].width = 18
    about.column_dimensions["B"].width = 90

    def sheet(title: str, columns, rows, widths, locked=()):
        ws = book.create_sheet(title)
        ws.append([label for _key, label in columns])
        for cell in ws[1]:
            cell.font = head_font
            cell.fill = head_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 30
        ws.freeze_panes = "B2" if locked else "A2"
        for row in rows:
            ws.append([_excel(row.get(key), key) for key, _label in columns])
        for index, (key, _label) in enumerate(columns, start=1):
            letter = ws.cell(row=1, column=index).column_letter
            ws.column_dimensions[letter].width = widths.get(key, 14)
            if key in DATE_FIELDS:
                for cell in ws[letter][1:]:
                    cell.number_format = "dd/mm/yyyy"
            if key in locked:
                for cell in ws[letter][1:]:
                    cell.fill = grey
                    cell.font = grey_font
        return ws

    projects_ws = sheet("Projects", PROJECT_COLUMNS, projects,
                        {"number": 14, "name": 40, "status": 16, "budget_mm": 12,
                         "notes": 40})
    rule = DataValidation(type="list", formula1='"' + ",".join(statuses) + '"',
                          allow_blank=True)
    projects_ws.add_data_validation(rule)
    rule.add(f"C2:C{max(2, len(projects) + 200)}")

    rows = []
    for item in items:
        issues = item.get("issues") or []
        last = issues[-1] if issues else {}
        rows.append({
            "ref": item["row"], "job": item["project_number"],
            "deliverable": item["name"],
            **{k: last.get(k) for k in ISSUE_FIELDS},
            "planned": last.get("planned") or (None if issues else item.get("status_date")),
            "starts": starts.get(item["row"]),
            "people": ", ".join(item.get("people") or []),
        })
    subs_ws = sheet("Submissions", SUBMISSION_COLUMNS, rows,
                    {"ref": 7, "job": 12, "deliverable": 40, "rev": 9, "purpose": 9,
                     "planned": 13, "submitted": 13, "returned": 13, "code": 10,
                     "reason": 34, "starts": 13, "people": 26},
                    locked=("ref", "job", "deliverable"))
    codes = DataValidation(type="list", formula1='"A,B,C"', allow_blank=True)
    purposes = DataValidation(type="list", formula1='"' + ",".join(revisions.PURPOSES) + '"',
                              allow_blank=True)
    subs_ws.add_data_validation(codes)
    subs_ws.add_data_validation(purposes)
    last_row = max(2, len(rows) + 1)
    codes.add(f"I2:I{last_row}")
    purposes.add(f"E2:E{last_row}")

    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _excel(value: Any, key: str) -> Any:
    if value in (None, ""):
        return None
    if key in DATE_FIELDS:
        day = as_date(value)
        return day if day else value
    return value


# --------------------------------------------------------------------------
# back in
# --------------------------------------------------------------------------

def _headers(row: Sequence[Any]) -> Dict[str, int]:
    return {as_text(v).lower(): i for i, v in enumerate(row) if as_text(v)}


def is_plan_book(filename: str, data: bytes) -> bool:
    """Whether a file is this workbook: an .xlsx with both sheets and the
    Ref column, whatever it was renamed to."""
    if not filename.lower().endswith((".xlsx", ".xlsm")):
        return False
    try:
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        return False
    try:
        if "Submissions" not in book.sheetnames or "Projects" not in book.sheetnames:
            return False
        first = next(book["Submissions"].iter_rows(max_row=1, values_only=True), ())
        heads = _headers(first)
        return "ref" in heads and "planned to go" in heads
    finally:
        book.close()


def read(filename: str, data: bytes) -> Dict[str, Any]:
    """Each row of both sheets, as typed, keyed by what it is."""
    from openpyxl import load_workbook
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as error:
        raise PlanBookError(f"{filename}: could not be opened ({error}).")
    try:
        def rows_of(name: str, columns) -> List[Tuple[int, Dict[str, Any]]]:
            grid = list(book[name].iter_rows(values_only=True))
            if not grid:
                return []
            heads = _headers(grid[0])
            where = {key: heads.get(label.lower()) for key, label in columns}
            out = []
            for number, row in enumerate(grid[1:], start=2):
                if not any(v not in (None, "") for v in row):
                    continue
                out.append((number, {key: (row[i] if i is not None and i < len(row) else None)
                                     for key, i in where.items()}))
            return out
        return {"filename": filename,
                "projects": rows_of("Projects", PROJECT_COLUMNS),
                "submissions": rows_of("Submissions", SUBMISSION_COLUMNS)}
    finally:
        book.close()


def _same(old: Any, new: Any) -> bool:
    return (old in (None, "") and new in (None, "")) or old == new


def _date_or_error(value: Any, label: str, where: str, errors: List[str]) -> Tuple[bool, Optional[str]]:
    if value in (None, ""):
        return True, None
    day = as_date(value)
    if day is None:
        errors.append(f"{where}: {label} {as_text(value)!r} is not a date. Write it day first, 25/12/2026.")
        return False, None
    return True, day.isoformat()


def changes(book: Dict[str, Any], *, projects: Sequence[Dict[str, Any]],
            items: Sequence[Dict[str, Any]], starts: Dict[int, Optional[str]],
            engineers: Sequence[str], statuses: Sequence[str]) -> Dict[str, Any]:
    """What bringing the workbook back would change, line by line, and what
    in it cannot be read.  Writes nothing."""
    out: List[Dict[str, Any]] = []
    errors: List[str] = []
    have = {p["number"]: p for p in projects}
    seen_numbers = set()

    for line, row in book["projects"]:
        where = f"Projects, row {line}"
        number = as_text(row.get("number"))
        if not number:
            errors.append(f"{where}: no job number, so it was left out.")
            continue
        if number in seen_numbers:
            errors.append(f"{where}: {number} is on the sheet twice; only the first is read.")
            continue
        seen_numbers.add(number)
        typed: Dict[str, Any] = {}
        bad = False
        for key in PROJECT_FIELDS:
            value = row.get(key)
            if key in ("start", "end"):
                ok, value = _date_or_error(value, LABELS[key], where, errors)
                bad = bad or not ok
            elif key == "budget_mm":
                if value not in (None, ""):
                    number_value = as_number(value)
                    if number_value is None or number_value <= 0:
                        errors.append(f"{where}: the budget {as_text(value)!r} is not a number above 0.")
                        bad = True
                    value = number_value
            else:
                value = as_text(value)
            typed[key] = value
        if typed["status"] and typed["status"] not in statuses:
            errors.append(f"{where}: the status {typed['status']!r} is not one of "
                          f"{', '.join(statuses)}.")
            bad = True
        if bad:
            continue
        current = have.get(number)
        if current is None:
            if not typed["name"] or not typed["budget_mm"]:
                errors.append(f"{where}: {number} is new; give it a name and a budget to add it.")
                continue
            out.append({"kind": "new_project", "key": number,
                        "what": f"{number} {typed['name']}",
                        "fields": [{"field": k, "label": LABELS[k], "old": None, "new": typed[k]}
                                   for k in PROJECT_FIELDS if typed[k] not in (None, "")],
                        "body": {"number": number, **typed,
                                 "status": typed["status"] or "Not Started"}})
            continue
        fields = []
        for key in PROJECT_FIELDS:
            old = current.get(key)
            new = typed[key]
            if key == "budget_mm" and new is None:
                continue                  # a budget is never cleared from Excel
            if key == "status" and not new:
                continue
            if key == "budget_mm" and old is not None and new is not None \
                    and abs(float(old) - float(new)) < 1e-9:
                continue
            if not _same(old, new):
                fields.append({"field": key, "label": LABELS[key], "old": old, "new": new})
        if fields:
            out.append({"kind": "project", "key": number,
                        "what": f"{number} {current.get('name') or ''}".strip(),
                        "fields": fields,
                        "body": {f["field"]: f["new"] for f in fields}})

    by_row = {item["row"]: item for item in items}
    seen_refs = set()
    for line, row in book["submissions"]:
        where = f"Submissions, row {line}"
        ref = as_number(row.get("ref"))
        if ref is None or int(ref) not in by_row:
            name = as_text(row.get("deliverable")) or "this row"
            errors.append(f"{where}: {name} is not a submission in the app (its Ref is "
                          f"{as_text(row.get('ref')) or 'empty'}). Add deliverables on the project first.")
            continue
        ref = int(ref)
        if ref in seen_refs:
            errors.append(f"{where}: Ref {ref} is on the sheet twice; only the first is read.")
            continue
        seen_refs.add(ref)
        item = by_row[ref]
        issues = item.get("issues") or []
        last = issues[-1] if issues else None
        typed = {}
        bad = False
        for key in ISSUE_FIELDS + ("starts",):
            value = row.get(key)
            if key in DATE_FIELDS:
                ok, value = _date_or_error(value, LABELS[key], where, errors)
                bad = bad or not ok
            elif key == "code":
                value = revisions.clean_code(value)
                if value == "?":
                    errors.append(f"{where}: the client code {as_text(row.get('code'))!r} "
                                  "is not A, B or C.")
                    bad = True
            else:
                value = as_text(value)
            typed[key] = value
        people_text = as_text(row.get("people"))
        people = [n.strip() for n in people_text.replace(";", ",").split(",") if n.strip()]
        unknown = [n for n in people if n not in engineers]
        if unknown:
            errors.append(f"{where}: {', '.join(unknown)} {'is' if len(unknown) == 1 else 'are'} "
                          "not on the team. Use the names as they are on Team.")
            bad = True
        if bad:
            continue
        what = f"{item['project_number']} {item['name']}"
        fields = []
        body: Dict[str, Any] = {}
        if last is None:
            for key in ISSUE_FIELDS:
                if typed[key] not in (None, "") and not (
                        key == "planned" and typed[key] == item.get("status_date")):
                    fields.append({"field": key, "label": LABELS[key], "old": None,
                                   "new": typed[key]})
                if typed[key] not in (None, ""):
                    body[key] = typed[key]
            start = bool(fields)
        else:
            for key in ISSUE_FIELDS:
                old = last.get(key)
                if not _same(old, typed[key]):
                    fields.append({"field": key, "label": LABELS[key], "old": old,
                                   "new": typed[key]})
                    body[key] = typed[key] if typed[key] is not None else ""
            start = False
        prep: Dict[str, Any] = {}
        if not _same(starts.get(ref), typed["starts"]):
            fields.append({"field": "starts", "label": LABELS["starts"],
                           "old": starts.get(ref), "new": typed["starts"]})
            prep["starts"] = typed["starts"] or ""
        if people and sorted(people) != sorted(item.get("people") or []):
            fields.append({"field": "people", "label": LABELS["people"],
                           "old": ", ".join(item.get("people") or []) or None,
                           "new": ", ".join(people)})
            prep["people"] = people
        if fields:
            out.append({"kind": "submission", "key": str(ref), "row": ref, "what": what,
                        "issue_id": last["id"] if last else None, "start": start,
                        "fields": fields, "body": body, "prep": prep})
    return {"changes": out, "errors": errors,
            "projects_read": len(book["projects"]),
            "submissions_read": len(book["submissions"])}
