"""The drawing list: drawing counts, issues and review codes from one file.

Every office keeps a drawing list -- a row a drawing, with its number, title,
revision, when it went to the client and the code it came back with.  It is
the honest source of the numbers the app wants, so rather than having
somebody count drawings and type a total, the list is uploaded:

* **how many drawings** each deliverable has is the number of rows for it;
* **how many are done** is how many have gone to the client;
* **what came back** -- codes A, B and C (or 1 to 4) -- says whether a
  deliverable is approved, approved with comments, or to be revised.

From those the app proposes what the register should say -- a deliverable
whose drawings have all gone is sent, with the date of the last; one whose
drawings all came back code A is accepted -- and one tap applies them, so the
submissions plan and the progress figures stay true without retyping.

Any drawing list with recognisable headings is read; the template is there
for a team that does not keep one yet.  A row is matched to a deliverable by
its job number and the deliverable's name (or its timesheet phase number);
a project with a single deliverable needs only the job number.
"""

from __future__ import annotations

import datetime as _dt
import io
import re
from collections import defaultdict
from typing import Any, Dict, Iterable, List, Mapping, Optional, Sequence, Tuple

from .model import ValidationError

#: The template's columns, in order, and what each may be headed in a list
#: somebody already keeps.
COLUMNS: Sequence[Tuple[str, str, Sequence[str]]] = (
    ("job_number", "Job Number", ("job number", "job no", "job", "project number",
                                  "project no", "project code", "jobnumber")),
    ("deliverable", "Deliverable", ("deliverable", "phase", "package",
                                    "deliverable description", "work package")),
    ("number", "Drawing No.", ("drawing no", "drawing number", "dwg no", "dwg",
                               "drawing", "document no", "document number", "doc no")),
    ("title", "Title", ("title", "drawing title", "description", "document title")),
    ("revision", "Revision", ("revision", "rev", "rev no")),
    ("status", "Status", ("status", "issue status", "purpose", "purpose of issue")),
    ("issued", "Issued", ("issued", "issue date", "date issued", "date of issue",
                          "transmitted", "submitted", "date submitted", "sent")),
    ("code", "Code", ("code", "return code", "review code", "client code",
                      "approval code", "response code")),
    ("returned", "Returned", ("returned", "return date", "date returned",
                              "comments received", "reply date", "response date")),
)
#: A status that means the drawing has gone to the client.
SENT = re.compile(r"\bIF[ACT]\b|for approval|for construction|for tender|issued|"
                  r"transmitted|submitted|\bsent\b", re.IGNORECASE)
#: A status that says the drawing has not gone yet, though it names issuing.
NOT_SENT = re.compile(r"\b(not|to be|un)[\s-]*(yet[\s-]+)?(issued|transmitted|submitted|sent)\b",
                      re.IGNORECASE)
#: Client codes, as letters; offices that number them use 1 to 4.
CODES = {"A": "A", "1": "A", "B": "B", "2": "B", "C": "C", "3": "C", "4": "C", "D": "C"}
#: Rules of Credit steps that mean "sent" and "accepted", by their names.
SENT_STEP = re.compile(r"transmit|issued|submitted|\bsent\b", re.IGNORECASE)
ACCEPTED_STEP = re.compile(r"accept|code 1|approv|comments closed|close ?out|final",
                           re.IGNORECASE)
#: The first rows a heading row is looked for in.
HEADER_SEARCH_ROWS = 20


class DrawingListError(ValidationError):
    pass


def _norm(text: Any) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(text or "").lower()).strip()


def _job(text: Any) -> str:
    return re.sub(r"\s+", "", str(text or "")).upper()


def _as_date(value: Any) -> Optional[_dt.date]:
    if isinstance(value, _dt.datetime):
        return value.date()
    if isinstance(value, _dt.date):
        return value
    text = str(value or "").strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d %b %Y",
                "%d-%b-%Y", "%d-%b-%y", "%d/%m/%y"):
        try:
            return _dt.datetime.strptime(text[:11].strip(), fmt).date()
        except ValueError:
            continue
    return None


# --------------------------------------------------------------------------
# the template
# --------------------------------------------------------------------------

def template(deliverables: Sequence[Dict[str, Any]] = ()) -> bytes:
    """An empty drawing list, with a starter row for each live deliverable."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Drawing list"
    sheet.append([label for _key, label, _aliases in COLUMNS])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1C5FA8")
    for item in deliverables:
        sheet.append([item.get("project_number"), item.get("name")])
    widths = (14, 34, 18, 40, 9, 22, 12, 8, 12)
    for index, width in enumerate(widths):
        sheet.column_dimensions[chr(ord("A") + index)].width = width
    sheet.freeze_panes = "A2"
    codes = DataValidation(type="list", formula1='"A,B,C"', allow_blank=True)
    status = DataValidation(type="list", allow_blank=True,
                            formula1='"In progress,For internal check,IFA - for approval,'
                                     'IFC - for construction,IFT - for tender,Superseded"')
    sheet.add_data_validation(codes)
    sheet.add_data_validation(status)
    codes.add("H2:H5000")
    status.add("F2:F5000")
    for column in ("G", "I"):
        for row in range(2, 2001):
            sheet[f"{column}{row}"].number_format = "yyyy-mm-dd"

    notes = book.create_sheet("How to fill")
    for line in (
        ["One row a drawing. Upload the file on Projects > Drawing list."],
        [],
        ["Job Number", "The job number, as on the timesheets."],
        ["Deliverable", "The deliverable's name as on Projects, or its timesheet "
                        "phase number. A project with one deliverable can leave it blank."],
        ["Drawing No., Title", "Either is enough for the row to count as a drawing."],
        ["Status", "IFA, IFC or IFT (or 'issued') means it has gone to the client."],
        ["Issued", "The date it went to the client. A date here also means it has gone."],
        ["Code", "What came back: A approved, B approved with comments, "
                 "C revise and resubmit (1 to 4 also work)."],
        ["Returned", "The date the code came back."],
        [],
        ["Rows with neither a drawing number nor a title are skipped, so the "
         "starter rows can stay until they are filled in."],
    ):
        notes.append(line)
    notes.column_dimensions["A"].width = 22
    notes.column_dimensions["B"].width = 90
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


# --------------------------------------------------------------------------
# reading a list
# --------------------------------------------------------------------------

def _header_map(cells: Sequence[Any]) -> Dict[str, int]:
    found: Dict[str, int] = {}
    names = [_norm(c) for c in cells]
    for key, label, aliases in COLUMNS:
        wanted = [_norm(label)] + [_norm(a) for a in aliases]
        for index, name in enumerate(names):
            if name and name in wanted and index not in found.values():
                found[key] = index
                break
    return found


def read(data: bytes, filename: str = "") -> List[Dict[str, Any]]:
    """Every drawing in a list, as plain values."""
    import openpyxl
    try:
        book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise DrawingListError([f"{filename or 'That file'} is not an Excel workbook."])
    try:
        sheets = sorted(book.worksheets,
                        key=lambda s: ("drawing" not in s.title.lower(), s.title))
        for sheet in sheets:
            # Some tools write a stub dimension that read-only mode believes.
            reset = getattr(sheet, "reset_dimensions", None)
            if reset is not None:
                reset()
            rows = list(sheet.iter_rows(values_only=True))
            for at, cells in enumerate(rows[:HEADER_SEARCH_ROWS]):
                columns = _header_map(cells or ())
                if "job_number" in columns and ("number" in columns or "title" in columns):
                    return _rows(rows[at + 1:], columns)
    finally:
        book.close()
    raise DrawingListError([
        f"No drawing list found in {filename or 'that file'}: it needs a heading row "
        f"with Job Number and Drawing No. (or Title). Download the template to see one."])


def _rows(rows: Iterable[Sequence[Any]], columns: Mapping[str, int]) -> List[Dict[str, Any]]:
    def get(cells, key):
        index = columns.get(key)
        return cells[index] if index is not None and index < len(cells) else None

    out = []
    for cells in rows:
        cells = cells or ()
        number = str(get(cells, "number") or "").strip()
        title = str(get(cells, "title") or "").strip()
        job = _job(get(cells, "job_number"))
        if not job or not (number or title):
            continue
        status = str(get(cells, "status") or "").strip()
        issued = _as_date(get(cells, "issued"))
        code_raw = str(get(cells, "code") or "").strip().upper()
        code = CODES.get(code_raw[-1:] if code_raw.startswith("CODE") else code_raw[:1])
        out.append({
            "job_number": job,
            "deliverable": get(cells, "deliverable"),
            "number": number, "title": title,
            "status": status,
            "issued": issued,
            "sent": bool(issued or (SENT.search(status) and not NOT_SENT.search(status)))
                    and "supersed" not in status.lower(),
            "code": code if code_raw else None,
            "returned": _as_date(get(cells, "returned")),
            "superseded": "supersed" in status.lower(),
        })
    return out


# --------------------------------------------------------------------------
# matching drawings to deliverables
# --------------------------------------------------------------------------

def _words(text: str) -> set:
    return {w for w in _norm(text).split() if len(w) > 1}


def _match_one(cell: Any, candidates: Sequence[Any]) -> Optional[Any]:
    if len(candidates) == 1 and not str(cell or "").strip():
        return candidates[0]
    text = str(cell or "").strip()
    if not text:
        return None
    if re.fullmatch(r"\d+(\.0+)?", text):
        phase = int(float(text))
        hits = [d for d in candidates if d.ts_phase == phase]
        if len(hits) == 1:
            return hits[0]
    wanted = _norm(text)
    exact = [d for d in candidates if _norm(d.name) == wanted]
    if len(exact) == 1:
        return exact[0]
    inside = [d for d in candidates
              if wanted and (wanted in _norm(d.name) or _norm(d.name) in wanted)]
    if len(inside) == 1:
        return inside[0]
    words = _words(text)
    best, score = None, 0.0
    for d in candidates:
        theirs = _words(d.name)
        if not words or not theirs:
            continue
        overlap = len(words & theirs) / len(words | theirs)
        if overlap > score:
            best, score = d, overlap
    if score >= 0.5:
        return best
    return candidates[0] if len(candidates) == 1 else None


def match(drawings: Sequence[Dict[str, Any]], deliverables: Sequence[Any]
          ) -> Dict[str, Any]:
    """Each drawing onto its deliverable, and a summary per deliverable."""
    by_project: Dict[str, List[Any]] = defaultdict(list)
    for d in deliverables:
        by_project[_job(d.project_number)].append(d)
    per: Dict[int, Dict[str, Any]] = {}
    unmatched: Dict[Tuple[str, str], int] = defaultdict(int)
    for drawing in drawings:
        if drawing["superseded"]:
            continue
        candidates = by_project.get(drawing["job_number"], [])
        target = _match_one(drawing["deliverable"], candidates) if candidates else None
        if target is None:
            unmatched[(drawing["job_number"], str(drawing["deliverable"] or ""))] += 1
            continue
        entry = per.setdefault(target.row, {
            "row": target.row, "project_number": target.project_number,
            "name": target.name, "total": 0, "issued": 0,
            "code_a": 0, "code_b": 0, "code_c": 0,
            "last_issued": None, "last_returned": None})
        entry["total"] += 1
        if drawing["sent"]:
            entry["issued"] += 1
            if drawing["issued"] and (entry["last_issued"] is None
                                      or drawing["issued"].isoformat() > entry["last_issued"]):
                entry["last_issued"] = drawing["issued"].isoformat()
        if drawing["code"]:
            entry[f"code_{drawing['code'].lower()}"] += 1
            if drawing["returned"] and (entry["last_returned"] is None
                                        or drawing["returned"].isoformat() > entry["last_returned"]):
                entry["last_returned"] = drawing["returned"].isoformat()
    return {
        "deliverables": sorted(per.values(), key=lambda e: (e["project_number"], e["row"])),
        "unmatched": [{"job_number": job, "deliverable": name, "drawings": count}
                      for (job, name), count in sorted(unmatched.items())],
        "drawings": sum(e["total"] for e in per.values()),
        "projects": sorted({_job(e["project_number"]) for e in per.values()}),
    }


# --------------------------------------------------------------------------
# what the register should say
# --------------------------------------------------------------------------

def _step(steps: Sequence[Any], type_code: str, pattern: "re.Pattern[str]",
          full: bool) -> Optional[Any]:
    mine = sorted((s for s in steps if s.type_code == type_code),
                  key=lambda s: (s.credit or 0, s.step_no))
    hits = [s for s in mine if pattern.search(s.step_name or "")
            and "internal" not in (s.step_name or "").lower()
            and ((s.credit or 0) >= 1.0 - 1e-9) == full]
    if hits:
        # The furthest such step: "issued to the client", not an internal issue.
        return hits[-1]
    if full:
        last = [s for s in mine if (s.credit or 0) >= 1.0 - 1e-9]
        return last[-1] if last else None
    return None


def proposals(summaries: Sequence[Dict[str, Any]], deliverables: Sequence[Any],
              steps: Sequence[Any], today: _dt.date) -> List[Dict[str, Any]]:
    """What each deliverable's register entry should say, by its drawings."""
    by_row = {d.row: d for d in deliverables}
    out = []
    for summary in summaries:
        d = by_row.get(summary["row"])
        if d is None or not summary["total"]:
            continue
        change: Dict[str, Any] = {}
        why: List[str] = []
        returned = summary["code_a"] + summary["code_b"] + summary["code_c"]
        all_sent = summary["issued"] >= summary["total"]
        if all_sent and not d.submitted_to_client:
            change["submitted_to_client"] = summary["last_issued"] or today.isoformat()
            why.append(f"all {summary['total']} drawings have gone to the client")
            sent = _step(steps, d.type_code, SENT_STEP, full=False)
            if sent and (d.step_no or 0) < sent.step_no:
                change["step_no"] = sent.step_no
                change["step_name"] = sent.step_name
        if returned and not d.comments_received and summary["last_returned"]:
            change["comments_received"] = summary["last_returned"]
            why.append(f"{returned} came back")
        if summary["code_a"] >= summary["total"] and not d.completed:
            change["completed"] = summary["last_returned"] or today.isoformat()
            why.append("every one is code A")
            accepted = _step(steps, d.type_code, ACCEPTED_STEP, full=True)
            if accepted and (d.step_no or 0) < accepted.step_no:
                change["step_no"] = accepted.step_no
                change["step_name"] = accepted.step_name
        if change:
            out.append({"row": d.row, "project_number": d.project_number,
                        "name": d.name, "change": change,
                        "why": "; ".join(why).capitalize() + "."})
    return out
