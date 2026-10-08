"""Round 3 bug fixes in the unit's storage: SQL, copies, migration, readers.

Each test fails on the code before its fix.  Synthetic data only.
"""

import json
import re
import sqlite3
import zipfile

import openpyxl
import pytest

from workload_app import config as cfg, legacy, storage
from workload_app.model import ValidationError, as_date, as_number, as_text, stored_date
from workload_app.unit import Unit
from workload_app.xlsx_io import Workbook


def _as_excel_saves_it(path):
    """Move every text cell into shared strings, the way Excel saves a file
    (openpyxl writes them inline)."""
    with zipfile.ZipFile(path) as zf:
        entries = {name: zf.read(name) for name in zf.namelist()}
    strings = []

    def share(match):
        strings.append(match.group(2))
        return f'<c r="{match.group(1)}" t="s"><v>{len(strings) - 1}</v></c>'

    for name in [n for n in entries if n.startswith("xl/worksheets/")]:
        entries[name] = re.sub(
            r'<c r="([A-Z]+\d+)" t="inlineStr"><is><t>(.*?)</t></is></c>',
            share, entries[name].decode("utf-8")).encode("utf-8")
    entries["xl/sharedStrings.xml"] = (
        '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        + "".join(f"<si><t>{text}</t></si>" for text in strings)
        + "</sst>").encode("utf-8")
    with zipfile.ZipFile(path, "w") as zf:
        for name, data in entries.items():
            zf.writestr(name, data)
    assert strings, "nothing was moved into shared strings"
    return path


def _tasks_workbook(path, odd=True):
    """A workbook with a Tasks sheet as Excel saves it: text as shared strings."""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = cfg.SHEET_TASKS
    cols = cfg.TASK_COLUMNS
    sheet[cfg.TASKS_SETTINGS_CELL] = json.dumps({"day_start": "08:00"})
    first = cfg.TASKS_FIRST_ROW
    sheet[f"{cols['id']}{first}"] = 7
    sheet[f"{cols['name']}{first}"] = "Check pile caps"
    sheet[f"{cols['assignees']}{first}"] = "Ahmed, Osama"
    sheet[f"{cols['status']}{first}"] = "In progress"
    sheet[f"{cols['deliverable_row']}{first}"] = 12
    if not odd:
        book.save(path)
        return _as_excel_saves_it(path)
    # An id typed as text, a row copied in Excel (same id) and a stray note.
    sheet[f"{cols['id']}{first + 1}"] = "8"
    sheet[f"{cols['name']}{first + 1}"] = "Draft sections"
    sheet[f"{cols['id']}{first + 2}"] = 7
    sheet[f"{cols['name']}{first + 2}"] = "Check pile caps (copy)"
    sheet[f"{cols['id']}{first + 3}"] = "see notes"
    sheet[f"{cols['name']}{first + 3}"] = "Not a task"
    book.save(path)
    return _as_excel_saves_it(path)


class TestOldTasksSheet:
    def test_text_saved_by_excel_reads_as_text(self, tmp_path):
        """Excel keeps text as shared strings; the migration read their
        index numbers, so every task came across named "0", "1"..."""
        tasks, settings = legacy._tasks(
            Workbook(_tasks_workbook(tmp_path / "t.xlsx", odd=False)))
        first = tasks[0]
        assert first.id == 7
        assert first.name == "Check pile caps"
        assert first.assignees == ["Ahmed", "Osama"]
        assert first.status == "In progress"
        assert first.deliverable_row == 12
        assert settings == {"day_start": "08:00"}

    def test_odd_ids_do_not_stop_the_unit_opening(self, tmp_path):
        """A text id, a repeated id or a note in the id column used to crash
        the move (or break the tasks table's unique id), so the unit could
        never be opened."""
        tasks, _ = legacy._tasks(Workbook(_tasks_workbook(tmp_path / "t.xlsx")))
        assert [(t.id, t.name) for t in tasks] == [
            (7, "Check pile caps"), (8, "Draft sections")]


class TestNumbersAndDatesNobodyMeans:
    def test_a_compact_date_is_refused_not_a_crash(self, wb):
        """"20261008" was read as an Excel serial past year 9999: a 500."""
        assert as_date("20261008") is None
        assert stored_date(1e12) is None
        with pytest.raises(ValidationError) as caught:
            wb.add_project({"number": "P-TEST-1", "name": "Quay wall",
                            "budget_mm": 2, "start": "20261008"})
        assert any("Start date" in e for e in caught.value.errors)

    def test_infinite_numbers_are_refused_not_a_crash(self, wb):
        assert as_number("inf") is None
        assert as_number(float("nan")) is None
        assert as_number("1e999") is None
        assert as_text(float("inf")) == "inf"
        with pytest.raises(ValidationError):
            wb.save_settings({"plan_year": "inf"})
        with pytest.raises(ValidationError):
            wb.save_settings({"availability_years": [1e999]})
        with pytest.raises(ValidationError):
            wb.add_project({"number": "P-TEST-2", "name": "Jetty",
                            "budget_mm": "nan"})

    def test_a_workbook_date_past_any_calendar_reads_as_none(self, tmp_path):
        book = openpyxl.Workbook()
        book.active.title = "Inputs"
        book.active["A1"] = 1e10
        book.active["A2"] = 46000
        book.save(tmp_path / "d.xlsx")
        read = Workbook(tmp_path / "d.xlsx")
        assert read.get_date("Inputs", "A1") is None
        assert read.get_date("Inputs", "A2").year == 2025


class TestRenameKeepsTheLead:
    def test_a_renamed_lead_still_leads_their_team(self, wb):
        """teams.lead names a person; a rename left the old name there, so
        the team showed a lead nobody is and pushed work to the real one."""
        wb.store.add_team("team0001", "Piles", "Osama")
        osama = next(e for e in wb.engineers() if e.short_name == "Osama")
        wb.update_engineer("Osama", {
            "short_name": "Osama A", "pattern": osama.pattern,
            "available_hours": osama.available_hours,
            "availability": osama.availability})
        assert [t["lead"] for t in wb.store.teams()] == ["Osama A"]


class TestDeletingAUnit:
    def test_the_last_state_is_kept_as_a_copy(self, tmp_path):
        """Deleting a unit promises its copies can be put back, but the
        newest copy could be half a day old: what was done since was lost."""
        path = storage.new_unit(tmp_path, 1, "unit0001")
        storage.keep_a_copy(tmp_path, 1, path)           # the copy on open
        Unit(path).add_engineer({"short_name": "Kirolos", "available_hours": 160})
        storage.remove_unit_file(tmp_path, 1, path.name)
        assert not path.exists()
        newest = storage.latest_backup(tmp_path, 1, "unit0001")
        db = sqlite3.connect(newest)
        try:
            names = [r[0] for r in db.execute("SELECT name FROM engineers")]
        finally:
            db.close()
        assert "Kirolos" in names


class TestRenameKeepsSavedWhatIfs:
    def test_a_saved_what_if_follows_a_renamed_person(self, wb):
        """A what-if keeps the names work moves from and to; a rename left
        the old name there, so the saved what-if went stale."""
        moves = [{"kind": "project", "project": "P-TEST-9", "from": "Osama",
                  "to": "Kirolos", "share": 0.5},
                 {"kind": "extra", "project": "Survey", "to": "Osama", "hours": 4}]
        wb.store.add_what_if("Lighten Kirolos", 10, json.dumps(moves))
        osama = next(e for e in wb.engineers() if e.short_name == "Osama")
        wb.update_engineer("Osama", {
            "short_name": "Osama A", "pattern": osama.pattern,
            "available_hours": osama.available_hours,
            "availability": osama.availability})
        kept = json.loads(wb.store.what_ifs()[0]["moves"])
        assert [m.get("from") for m in kept] == ["Osama A", None]
        assert [m["to"] for m in kept] == ["Kirolos", "Osama A"]
